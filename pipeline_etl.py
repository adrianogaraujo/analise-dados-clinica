"""
pipeline_etl.py
Pipeline de Análise de Dados — Clínica Médica Exemplo
Projeto de portfólio | Sidia — Estágio Análise de Dados e Automação

Fluxo:
    1. Ingestão dos CSVs (atendimentos + médicos)
    2. Limpeza e transformação
    3. Cálculo de KPIs
    4. Exportação de relatório Excel formatado

Dependências: pandas, openpyxl
    pip install pandas openpyxl
"""

import pandas as pd
import os
from datetime import datetime

# ─────────────────────────────────────────────
# 1. INGESTÃO
# ─────────────────────────────────────────────

def carregar_dados(caminho_atendimentos: str, caminho_medicos: str) -> tuple:
    """Carrega os CSVs e retorna dois DataFrames."""
    df_at = pd.read_csv(caminho_atendimentos, parse_dates=["data_agendamento"])
    df_med = pd.read_csv(caminho_medicos)
    print(f"[✓] Atendimentos carregados: {len(df_at)} registros")
    print(f"[✓] Médicos carregados:      {len(df_med)} registros")
    return df_at, df_med


# ─────────────────────────────────────────────
# 2. LIMPEZA E TRANSFORMAÇÃO
# ─────────────────────────────────────────────

def limpar_e_transformar(df: pd.DataFrame) -> pd.DataFrame:
    """
    Aplica limpeza básica e cria colunas derivadas.
    - Remove duplicatas
    - Garante tipos corretos
    - Cria coluna 'trimestre' e 'realizado' (flag booleana)
    """
    # Remover duplicatas pelo ID
    antes = len(df)
    df = df.drop_duplicates(subset=["id_atendimento"])
    removidos = antes - len(df)
    if removidos > 0:
        print(f"[!] {removidos} duplicatas removidas")

    # Coluna de trimestre
    df["trimestre"] = df["data_agendamento"].dt.quarter.map(
        {1: "Q1", 2: "Q2", 3: "Q3", 4: "Q4"}
    )

    # Flag: atendimento efetivamente realizado
    df["realizado"] = df["status"] == "Realizado"

    # Garantir que valor seja numérico
    df["valor_cobrado"] = pd.to_numeric(df["valor_cobrado"], errors="coerce").fillna(0.0)

    print(f"[✓] Transformações aplicadas")
    return df


# ─────────────────────────────────────────────
# 3. CÁLCULO DE KPIs
# ─────────────────────────────────────────────

def calcular_kpis(df: pd.DataFrame) -> dict:
    """
    Retorna dicionário com os principais KPIs da clínica.
    """
    total_agendamentos = len(df)
    total_realizados = df["realizado"].sum()
    taxa_realizacao = total_realizados / total_agendamentos * 100
    taxa_cancelamento = (df["status"] == "Cancelado").sum() / total_agendamentos * 100
    taxa_falta = (df["status"] == "Faltou").sum() / total_agendamentos * 100
    receita_total = df["valor_cobrado"].sum()
    ticket_medio = df.loc[df["realizado"], "valor_cobrado"].mean()

    kpis = {
        "Total de Agendamentos": total_agendamentos,
        "Atendimentos Realizados": int(total_realizados),
        "Taxa de Realização (%)": round(taxa_realizacao, 1),
        "Taxa de Cancelamento (%)": round(taxa_cancelamento, 1),
        "Taxa de Falta (%)": round(taxa_falta, 1),
        "Receita Total (R$)": round(receita_total, 2),
        "Ticket Médio (R$)": round(ticket_medio, 2),
    }

    print("\n─── KPIs Gerais ───────────────────────────")
    for chave, valor in kpis.items():
        print(f"  {chave:<35} {valor}")
    print("────────────────────────────────────────────\n")

    return kpis


def agrupar_por_especialidade(df: pd.DataFrame) -> pd.DataFrame:
    """Atendimentos e receita por especialidade, apenas realizados."""
    return (
        df[df["realizado"]]
        .groupby("especialidade")
        .agg(
            atendimentos=("id_atendimento", "count"),
            receita_total=("valor_cobrado", "sum"),
            ticket_medio=("valor_cobrado", "mean"),
        )
        .round(2)
        .sort_values("receita_total", ascending=False)
        .reset_index()
    )


def agrupar_por_convenio(df: pd.DataFrame) -> pd.DataFrame:
    """Volume e receita por convênio."""
    return (
        df[df["realizado"]]
        .groupby("convenio")
        .agg(
            atendimentos=("id_atendimento", "count"),
            receita_total=("valor_cobrado", "sum"),
        )
        .round(2)
        .sort_values("atendimentos", ascending=False)
        .reset_index()
    )


def agrupar_por_mes(df: pd.DataFrame) -> pd.DataFrame:
    """Evolução mensal de atendimentos e receita."""
    return (
        df[df["realizado"]]
        .groupby("mes")
        .agg(
            atendimentos=("id_atendimento", "count"),
            receita=("valor_cobrado", "sum"),
        )
        .round(2)
        .reset_index()
        .rename(columns={"mes": "Mês"})
    )


def agrupar_por_medico(df: pd.DataFrame) -> pd.DataFrame:
    """Ranking de produtividade por médico."""
    return (
        df[df["realizado"]]
        .groupby("medico")
        .agg(
            atendimentos=("id_atendimento", "count"),
            receita_gerada=("valor_cobrado", "sum"),
            duracao_media_min=("duracao_minutos", "mean"),
        )
        .round(2)
        .sort_values("atendimentos", ascending=False)
        .reset_index()
    )


def tabela_status(df: pd.DataFrame) -> pd.DataFrame:
    """Distribuição de status de atendimento."""
    contagem = df["status"].value_counts().reset_index()
    contagem.columns = ["status", "quantidade"]
    contagem["percentual"] = (contagem["quantidade"] / len(df) * 100).round(1)
    return contagem


# ─────────────────────────────────────────────
# 4. EXPORTAÇÃO EXCEL FORMATADO
# ─────────────────────────────────────────────

def exportar_excel(
    kpis: dict,
    df_especialidade: pd.DataFrame,
    df_convenio: pd.DataFrame,
    df_mensal: pd.DataFrame,
    df_medico: pd.DataFrame,
    df_status: pd.DataFrame,
    caminho_saida: str,
):
    """
    Gera relatório Excel com múltiplas abas formatadas.
    Cada aba corresponde a uma visão analítica.
    """
    os.makedirs(os.path.dirname(caminho_saida), exist_ok=True)

    with pd.ExcelWriter(caminho_saida, engine="openpyxl") as writer:

        # ── Aba 1: KPIs ──────────────────────────────
        df_kpis = pd.DataFrame(
            list(kpis.items()), columns=["Indicador", "Valor"]
        )
        df_kpis.to_excel(writer, sheet_name="KPIs", index=False, startrow=1)

        ws_kpi = writer.sheets["KPIs"]
        ws_kpi["A1"] = "PAINEL DE INDICADORES — CLÍNICA MÉDICA"
        _aplicar_estilo_cabecalho(ws_kpi, linha=2, n_colunas=2)
        _ajustar_largura_colunas(ws_kpi)

        # ── Aba 2: Por Especialidade ──────────────────
        df_especialidade.to_excel(
            writer, sheet_name="Por Especialidade", index=False, startrow=1
        )
        ws_esp = writer.sheets["Por Especialidade"]
        ws_esp["A1"] = "ATENDIMENTOS E RECEITA POR ESPECIALIDADE"
        _aplicar_estilo_cabecalho(ws_esp, linha=2, n_colunas=len(df_especialidade.columns))
        _ajustar_largura_colunas(ws_esp)

        # ── Aba 3: Por Convênio ───────────────────────
        df_convenio.to_excel(
            writer, sheet_name="Por Convênio", index=False, startrow=1
        )
        ws_conv = writer.sheets["Por Convênio"]
        ws_conv["A1"] = "VOLUME E RECEITA POR CONVÊNIO"
        _aplicar_estilo_cabecalho(ws_conv, linha=2, n_colunas=len(df_convenio.columns))
        _ajustar_largura_colunas(ws_conv)

        # ── Aba 4: Evolução Mensal ────────────────────
        df_mensal.to_excel(
            writer, sheet_name="Evolução Mensal", index=False, startrow=1
        )
        ws_men = writer.sheets["Evolução Mensal"]
        ws_men["A1"] = "EVOLUÇÃO MENSAL DE ATENDIMENTOS E RECEITA"
        _aplicar_estilo_cabecalho(ws_men, linha=2, n_colunas=len(df_mensal.columns))
        _ajustar_largura_colunas(ws_men)

        # ── Aba 5: Ranking Médicos ────────────────────
        df_medico.to_excel(
            writer, sheet_name="Ranking Médicos", index=False, startrow=1
        )
        ws_med = writer.sheets["Ranking Médicos"]
        ws_med["A1"] = "PRODUTIVIDADE POR MÉDICO"
        _aplicar_estilo_cabecalho(ws_med, linha=2, n_colunas=len(df_medico.columns))
        _ajustar_largura_colunas(ws_med)

        # ── Aba 6: Status ─────────────────────────────
        df_status.to_excel(
            writer, sheet_name="Status Atendimentos", index=False, startrow=1
        )
        ws_st = writer.sheets["Status Atendimentos"]
        ws_st["A1"] = "DISTRIBUIÇÃO DE STATUS DOS AGENDAMENTOS"
        _aplicar_estilo_cabecalho(ws_st, linha=2, n_colunas=len(df_status.columns))
        _ajustar_largura_colunas(ws_st)

    print(f"[✓] Relatório Excel exportado → {caminho_saida}")


# ─────────────────────────────────────────────
# Funções auxiliares de formatação
# ─────────────────────────────────────────────

def _aplicar_estilo_cabecalho(ws, linha: int, n_colunas: int):
    """Aplica negrito e cor azul ao cabeçalho da tabela."""
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    fill_azul = PatternFill("solid", fgColor="1F4E79")
    fonte_branca = Font(bold=True, color="FFFFFF")

    for col in range(1, n_colunas + 1):
        cell = ws.cell(row=linha, column=col)
        cell.fill = fill_azul
        cell.font = fonte_branca
        cell.alignment = Alignment(horizontal="center")


def _ajustar_largura_colunas(ws):
    """Ajusta a largura de cada coluna ao conteúdo."""
    for col in ws.columns:
        max_len = max(
            (len(str(cell.value)) for cell in col if cell.value is not None),
            default=10
        )
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)


# ─────────────────────────────────────────────
# 5. EXECUÇÃO PRINCIPAL
# ─────────────────────────────────────────────

if __name__ == "__main__":
    print("=" * 50)
    print("  PIPELINE ETL — CLÍNICA MÉDICA")
    print(f"  Executado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    print("=" * 50 + "\n")

    # Caminhos
    ATENDIMENTOS = "data/atendimentos.csv"
    MEDICOS      = "data/medicos.csv"
    SAIDA_EXCEL  = "outputs/relatorio_clinica.xlsx"

    # Executar pipeline
    df_at, df_med = carregar_dados(ATENDIMENTOS, MEDICOS)
    df_at         = limpar_e_transformar(df_at)
    kpis          = calcular_kpis(df_at)

    df_esp   = agrupar_por_especialidade(df_at)
    df_conv  = agrupar_por_convenio(df_at)
    df_mes   = agrupar_por_mes(df_at)
    df_medpr = agrupar_por_medico(df_at)
    df_st    = tabela_status(df_at)

    exportar_excel(kpis, df_esp, df_conv, df_mes, df_medpr, df_st, SAIDA_EXCEL)

    print("\n[✓] Pipeline concluído com sucesso.")
